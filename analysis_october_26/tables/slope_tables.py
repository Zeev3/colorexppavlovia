# Slide-ready tables (editable PowerPoint): forward and reverse slope against 1 at each level, and the change between
# levels. Experience instruction only. Same layout as tracking_table.py.
# Input: output/slope_vs_one.csv, output/experience_only_results.csv.
# Output: output/tables/slopes_vs_one_table.pptx: forward Exp 4, forward Exps 1-3, reverse Exp 4, reverse Exps 1-3,
# then Exp 4 figure slides: ellipses overlaid, each level with both regression lines (together and one per level), slopes four ways.
import csv, os, sys
from decimal import Decimal, ROUND_HALF_UP
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
OUT = sys.argv[1]
INK, LIGHT, WHITE, MUTED = RGBColor(0x1E, 0x2A, 0x36), RGBColor(0xF3, 0xF5, 0xF8), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x55, 0x61, 0x6D)
S1 = list(csv.DictReader(open(os.path.join(OUT, "slope_vs_one.csv"))))
ER = list(csv.DictReader(open(os.path.join(OUT, "experience_only_results.csv"))))
def P(p):
    p = float(p); return "p < .001" if p < .001 else "p = " + (("%.3f" % p).lstrip("0") if p < .1 else ("%.2f" % p).lstrip("0"))
def r2(x, q="0.01"):
    s = str(Decimal(str(x)).quantize(Decimal(q), ROUND_HALF_UP)); return "−" + s[1:] if s.startswith("-") else s
def sg(x, q="0.01"):
    s = r2(x, q); return s if s.startswith("−") else "+" + s
def lab(e, l): return ("SD %s" if e == "exp4" else "%s sq.") % (str(l).replace(".0", ""))
NAME = {"exp1": "Exp 1", "exp2": "Exp 2", "exp3": "Exp 3", "exp4": "Exp 4"}
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
def text(s, x, y, w, h, t, size, bold=False, color=INK, font="Calibri"):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
def table(s, x, y, w, rows, colw, size=13, row_h=0.42):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = INK if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            c.margin_left = c.margin_right = Inches(0.07); r = c.text_frame.paragraphs[0].add_run(); r.text = v
            r.font.size = Pt(size); r.font.name = "Calibri"; r.font.bold = i == 0 or j == 0; r.font.color.rgb = WHITE if i == 0 else INK
def vs1_rows(exps, d, with_exp):
    rows = [(["Experiment"] if with_exp else []) + ["Level", "Mean slope [95% CI]", "People below 1", "vs 1: t (df), p (Holm)", "Model slope [95% CI], p vs 1"]]
    for e in exps:
        pp_ = [r for r in S1 if r["experiment"] == e and r["direction"] == d and r["source"] == "per-person t-test"]
        mm = {r["level"]: r for r in S1 if r["experiment"] == e and r["direction"] == d and r["source"] == "linked model"}
        for k, r in enumerate(pp_):
            m = mm[r["level"]]
            rows.append(([NAME[e] if k == 0 else ""] if with_exp else []) + [lab(e, r["level"]),
                "%s [%s, %s]" % (r2(r["slope"]), r2(r["ci_low"]), r2(r["ci_high"])), "%s of %s" % (r["n_below_1"], r["n"]),
                "%s (%d), %s" % (r2(r["t"], "0.1"), round(float(r["df"])), P(r["p_holm"])),
                "%s [%s, %s], %s" % (r2(m["slope"]), r2(m["ci_low"]), r2(m["ci_high"]), P(m["p_holm"]))])
    return rows
def pair_rows(e, meas):
    eff = next(r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: " + meas and r["test"] == "level effect")
    rows = [["Between levels (each person's slope)", "Difference", "t (df)", "p (Holm)"]]
    for r in [r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: " + meas and r["test"].startswith("pair")]:
        a, b = r["test"].replace("pair X", "").split(" - X"); est, tt = r["value"].split(", t")
        rows.append(["%s vs %s" % (lab(e, a), lab(e, b)), sg(est), ("t" + tt).replace("-", "−"), P(r["p"])])
    lin = next(r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: " + meas and r["test"] == "linear trend")
    rows.append(["Linear trend", "", lin["value"].split(", ")[1].replace("-", "−"), P(lin["p"])])
    return rows, eff["value"].split(", ges")[0] + ", " + P(eff["p"])
DIRS = [("forward", "slope_fwd", "Forward slope (report on true colour)", "Exp 4: forward slope against 1 at each level, and between levels"),
        ("reverse", "slope_rev", "Reverse slope (true colour on report)", "Exp 4: reverse slope against 1 at each level, and between levels")]
for d, meas, what, title in DIRS:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.35, 12.1, 0.6, title, 24, True, font="Cambria")
    meaning = ("Slope 1 = the report follows the colour exactly; below 1 = reports pulled toward the middle (regression to the mean)." if d == "forward" else
               "Slope 1 = the array behind a report is as extreme as the report; below 1 = arrays less extreme than the reports.")
    text(s, 0.6, 1.0, 12.1, 0.4, "%s, experience instruction, 50 participants. %s" % (what, meaning), 13, color=MUTED)
    table(s, 0.6, 1.55, 12.13, vs1_rows(["exp4"], d, False), [1.2, 2.6, 2.0, 3.0, 3.33], size=14, row_h=0.5)
    pr, eff = pair_rows("exp4", meas)
    table(s, 0.6, 3.8, 8.4, pr, [3.6, 1.5, 1.8, 1.5], size=14, row_h=0.48)
    text(s, 9.3, 3.8, 3.4, 2.6, "Level effect (repeated-measures ANOVA on each person's slope): %s." % eff, 14)
    text(s, 0.6, 6.55, 12.1, 0.6, "Mean slope and t-test: each person's own OLS slope at that level. Model: linked mixed model (its CIs adjusted for the three levels). Holm across the three levels.", 12, color=MUTED)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.35, 12.1, 0.6, "Exps 1–3: %s slope against 1 at every array size (experience)" % d, 24, True, font="Cambria")
    table(s, 0.6, 1.15, 12.13, vs1_rows(["exp1", "exp2", "exp3"], d, True), [1.1, 0.95, 2.35, 1.65, 2.85, 3.23], size=12, row_h=0.42)
    eff_txt = "; ".join("%s: %s" % (NAME[e], pair_rows(e, meas)[1]) for e in ("exp1", "exp2", "exp3"))
    text(s, 0.6, 6.0, 12.1, 1.0, "Level effect on each person's slope: " + eff_txt + ". Mean slope and t-test: each person's own OLS slope; model: linked mixed model (CIs adjusted for the number of sizes); Holm across sizes.", 12, color=MUTED)
# ---- figure slides (Exp 4, experience): ellipses and the slopes computed four ways
FIGD = os.path.join(OUT, "figures")
def pic_slide(title, img, note, w=None, h=None, x=0.6, y=1.25):
    s = prs.slides.add_slide(prs.slide_layouts[6]); text(s, 0.6, 0.35, 12.1, 0.6, title, 24, True, font="Cambria")
    kw = {"width": Inches(w)} if w else {"height": Inches(h)}
    s.shapes.add_picture(os.path.join(FIGD, img), Inches(x), Inches(y), **kw)
    text(s, 0.6, 6.85, 12.1, 0.5, note, 12, color=MUTED)
POOL = "Slopes and r on the figure: one regression over all trials, each person centred (overall across participants). The tested values are each person's own: forward 1.10 / 0.97 / 0.85, reverse 0.63 / 0.62 / 0.60."
pic_slide("Exp 4: tracking as data ellipses, all levels together (experience)", "draft_tracking_ellipses_exp4.png",
          "Each ellipse holds about 68% of a level's trials. Thinner = closer tracking (r); shorter = reports less spread out. " + POOL[:POOL.index(" The tested")], h=5.5, x=3.9)
pic_slide("Exp 4: each level, with both regression lines (experience)", "galton_ellipse_exp4.png",
          "Grey: every trial. Yellow: 68% ellipse. Red: forward line (report on true colour), through the leftmost and rightmost points. Blue: reverse line, through the lowest and highest points. " + POOL, w=11.4, x=0.95)
for l in ("3", "5", "7"):
    pic_slide("Exp 4, SD %s: trials, ellipse and both regression lines (experience)" % l, "galton_ellipse_exp4_SD%s.png" % l, POOL, h=5.5, x=4.0)
pic_slide("Exp 4: the same slopes computed four ways (experience)", "slopes_four_ways_exp4.png",
          "Forward: all four agree. Reverse: pooling trials across people lowers the slope (differences between people inflate the spread of the reports); the model and the per-person slopes are the ones to report.", w=11.0, x=1.15)

os.makedirs(os.path.join(OUT, "tables"), exist_ok=True)
prs.save(os.path.join(OUT, "tables", "slopes_vs_one_table.pptx")); print("saved")
