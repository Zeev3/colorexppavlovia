# Slide-ready tables (editable PowerPoint): tracking r against 1 at each level, and the change between levels.
# Experience instruction only. Input: output/r_vs_one.csv, output/experience_only_results.csv.
# Output: output/tables/tracking_r_table.pptx (slide 1: Exp 4; slide 2: Exps 1-3).
import csv, os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
OUT = sys.argv[1]
INK, LIGHT, WHITE, MUTED = RGBColor(0x1E, 0x2A, 0x36), RGBColor(0xF3, 0xF5, 0xF8), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x55, 0x61, 0x6D)
RV = list(csv.DictReader(open(os.path.join(OUT, "r_vs_one.csv"))))
ER = list(csv.DictReader(open(os.path.join(OUT, "experience_only_results.csv"))))
def P(p):
    p = float(p); return "p < .001" if p < .001 else "p = " + (("%.3f" % p).lstrip("0") if p < .1 else ("%.2f" % p).lstrip("0"))
def lab(e, l): return ("SD %s" if e == "exp4" else "%s sq.") % l
NAME = {"exp1": "Exp 1", "exp2": "Exp 2", "exp3": "Exp 3", "exp4": "Exp 4"}
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
def text(s, x, y, w, h, t, size, bold=False, color=INK, font="Calibri"):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = t; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
def table(s, x, y, w, rows, colw, size=13, row_h=0.42):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows))).table
    for j, cw in enumerate(colw): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(row):
            c = t.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = INK if i == 0 else (LIGHT if i % 2 == 0 else WHITE)
            c.margin_left = c.margin_right = Inches(0.07); r = c.text_frame.paragraphs[0].add_run(); r.text = v
            r.font.size = Pt(size); r.font.name = "Calibri"; r.font.bold = i == 0 or j == 0; r.font.color.rgb = WHITE if i == 0 else INK
def vs1_rows(exps, with_exp):
    rows = [(["Experiment"] if with_exp else []) + ["Level", "Mean r [95% CI]", "People with r < 1", "r vs 1: t (df)", "p (Holm)"]]
    for e in exps:
        for k, r in enumerate([r for r in RV if r["experiment"] == e]):
            rows.append(([NAME[e] if k == 0 else ""] if with_exp else []) + [lab(e, r["level"]), "%.2f [%.2f, %.2f]" % (float(r["mean_r"]), float(r["ci_low"]), float(r["ci_high"])),
                         "%s of %s" % (r["n_below_1"], r["n"]), ("%.1f (%d)" % (float(r["t"]), round(float(r["df"])))).replace("-", "−"), P(r["p_holm"])])
    return rows
def pair_rows(e):
    eff = next(r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: z" and r["test"] == "level effect")
    rows = [["Between levels (Fisher z)", "Difference", "t (df)", "p (Holm)"]]
    for r in [r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: z" and r["test"].startswith("pair")]:
        a, b = r["test"].replace("pair X", "").split(" - X"); est, tt = r["value"].split(", t")
        from decimal import Decimal, ROUND_HALF_UP
        d_ = str(Decimal(est).quantize(Decimal("0.01"), ROUND_HALF_UP)); d_ = "−" + d_[1:] if d_.startswith("-") else "+" + d_
        rows.append(["%s vs %s" % (lab(e, a), lab(e, b)), d_, ("t" + tt).replace("-", "−"), P(r["p"])])
    lin = next(r for r in ER if r["experiment"] == e and r["analysis"] == "RM ANOVA: z" and r["test"] == "linear trend")
    rows.append(["Linear trend", "", lin["value"].split(", ")[1].replace("-", "−"), P(lin["p"])])
    return rows, eff["value"].split(", ges")[0] + ", " + P(eff["p"])
# slide 1: Exp 4
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 0.6, 0.35, 12, 0.6, "Exp 4: tracking is below 1 at every level and falls with variance", 24, True, font="Cambria")
text(s, 0.6, 1.0, 12, 0.4, "Experience instruction, 50 participants. r = each person's correlation of report with true colour at that level (mean of the people's r).", 14, color=MUTED)
table(s, 0.6, 1.6, 12.13, vs1_rows(["exp4"], False), [1.6, 3.2, 2.6, 2.4, 2.33], size=15, row_h=0.5)
pr, eff = pair_rows("exp4")
table(s, 0.6, 3.85, 8.2, pr, [3.2, 1.6, 1.8, 1.6], size=15, row_h=0.5)
text(s, 9.1, 3.85, 3.6, 2.6, "Level effect (repeated-measures ANOVA on Fisher z): %s." % eff, 14)
text(s, 0.6, 6.6, 12.1, 0.6, "r vs 1 is descriptive: r cannot exceed 1. The test of interest is the change between levels.", 12, color=MUTED)
# slide 2: Exps 1-3
s = prs.slides.add_slide(prs.slide_layouts[6])
text(s, 0.6, 0.35, 12, 0.6, "Experiments 1–3: tracking against 1 at every array size (experience)", 24, True, font="Cambria")
table(s, 0.6, 1.2, 12.13, vs1_rows(["exp1", "exp2", "exp3"], True), [1.5, 1.3, 3.1, 2.3, 2.1, 1.83], size=13, row_h=0.42)
eff_txt = "; ".join("%s: %s" % (NAME[e], pair_rows(e)[1]) for e in ("exp1", "exp2", "exp3"))
text(s, 0.6, 6.15, 12.1, 0.9, "Level effect on r (Fisher z): " + eff_txt + ". r vs 1 is descriptive.", 12, color=MUTED)
os.makedirs(os.path.join(OUT, "tables"), exist_ok=True)
prs.save(os.path.join(OUT, "tables", "tracking_r_table.pptx")); print("saved")
